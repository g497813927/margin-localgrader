"""Generated fixtures cover text extraction without external documents or OCR."""

import codecs
import io
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image
from pypdf import PdfReader, PdfWriter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from grader import extract
from grader.extract import extract_document


def png_bytes():
    buffer = io.BytesIO()
    Image.new("RGB", (12, 12), color="red").save(buffer, "PNG")
    return buffer.getvalue()


def pdf_bytes(pages):
    """Each tuple is (text or None, include_image)."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    for text, image in pages:
        if text:
            pdf.drawString(40, 740, text)
        if image:
            pdf.drawImage(ImageReader(io.BytesIO(png_bytes())), 40, 600, 70, 70)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def docx_bytes(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def replace_zip_part(data, name, replacement):
    source = zipfile.ZipFile(io.BytesIO(data))
    target_bytes = io.BytesIO()
    with zipfile.ZipFile(target_bytes, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry.filename, replacement if entry.filename == name else source.read(entry.filename))
    return target_bytes.getvalue()


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16"])
def test_text_unicode_and_newline_normalization(encoding):
    result = extract_document("../essay.TXT", "你好 café\r\nsecond line".encode(encoding))
    assert result == {"filename": "essay.TXT", "text": "你好 café\nsecond line", "warnings": []}


def test_utf16_big_endian():
    result = extract_document("essay.txt", codecs.BOM_UTF16_BE + "Hello 世界".encode("utf-16-be"))
    assert result["text"] == "Hello 世界"


@pytest.mark.parametrize("data", [b"bad\xffutf8", b"binary\x00data", b"bad\x01control", codecs.BOM_UTF32_LE + b"a\0\0\0"])
def test_malformed_or_binary_text_rejected(data):
    with pytest.raises(ValueError, match="UTF-8|UTF-32"):
        extract_document("essay.txt", data)


def test_whitespace_has_empty_text_warning():
    result = extract_document("essay.txt", b"  \n ")
    assert result["text"] == ""
    assert "No extractable text" in result["warnings"][0]


@pytest.mark.parametrize("name,data,message", [
    ("essay.txt", b"", "empty"),
    ("essay.doc", b"old word", "save it as .docx"),
    ("essay.png", b"image", ".txt, .pdf, or .docx"),
    ("essay.pdf", b"not a PDF", "PDF could not be read"),
    ("essay.docx", b"not a ZIP", "Word document could not be read"),
])
def test_unsupported_or_corrupt_files(name, data, message):
    with pytest.raises(ValueError, match=message):
        extract_document(name, data)


def test_pdf_text_and_page_count():
    result = extract_document("essay.pdf", pdf_bytes([("First page text", False), ("Second page text", False)]))
    assert result["text"] == "First page text\n\nSecond page text"
    assert result["pages"] == 2


def test_pdf_image_only_has_no_text():
    result = extract_document("scan.pdf", pdf_bytes([(None, True)]))
    assert result["text"] == ""
    assert any("page(s): 1" in warning for warning in result["warnings"])
    assert any("No extractable text was found" in warning for warning in result["warnings"])
    assert set(result) == {"filename", "text", "warnings", "pages"}


def test_pdf_mixed_preserves_only_text():
    result = extract_document("mixed.pdf", pdf_bytes([("Typed answer", True), (None, True)]))
    assert result["text"] == "Typed answer"
    assert any("page(s): 2" in warning for warning in result["warnings"])


@pytest.mark.parametrize("password", ["secret", ""])
def test_encrypted_pdf_is_rejected(password):
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(pdf_bytes([("Private text", False)]))))
    writer.encrypt(password)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match="encrypted or password-protected"):
        extract_document("locked.pdf", buffer.getvalue())


def test_pdf_page_limit(monkeypatch):
    monkeypatch.setattr(extract, "MAX_PDF_PAGES", 1)
    with pytest.raises(ValueError, match="1-page limit"):
        extract_document("too-long.pdf", pdf_bytes([("One", False), ("Two", False)]))


def test_pdf_compressed_stream_limit_rejects_incomplete_extraction(monkeypatch):
    monkeypatch.setattr(extract, "MAX_PDF_STREAM_BYTES", 1024)
    with pytest.raises(ValueError, match="partial document was not accepted"):
        extract_document("compressed.pdf", pdf_bytes([("x" * 4096, False)]))


def test_pdf_extraction_failure_rejects_partial_document(monkeypatch):
    from pypdf._page import PageObject

    original = PageObject.extract_text

    def fail_on_first_page(self, *args, **kwargs):
        text = original(self, *args, **kwargs)
        if "bad page" in text:
            raise RuntimeError("corrupt font")
        return text

    monkeypatch.setattr(PageObject, "extract_text", fail_on_first_page)
    with pytest.raises(ValueError, match="partial document was not accepted"):
        extract_document("partial.pdf", pdf_bytes([("bad page", False), ("Good page", False)]))


def test_docx_paragraph_table_image_and_headers_in_order():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Course header"
    document.add_paragraph("Before the table")
    table = document.add_table(rows=2, cols=2)
    for cell, text in zip([*table.rows[0].cells, *table.rows[1].cells], ["Criterion", "Answer", "Evidence", "A quotation"]):
        cell.text = text
    document.add_picture(io.BytesIO(png_bytes()))
    document.add_paragraph("After the picture")
    document.sections[0].footer.paragraphs[0].text = "Course footer"
    result = extract_document("answer.docx", docx_bytes(document))
    assert result["text"] == "Course header\nBefore the table\nCriterion\tAnswer\nEvidence\tA quotation\nAfter the picture\nCourse footer"
    assert set(result) == {"filename", "text", "warnings"}


def test_docx_ignores_image_metadata_and_drawing_text():
    document = Document()
    paragraph = document.add_paragraph("Real response")
    run = paragraph.add_run()
    drawing = OxmlElement("w:drawing")
    text = OxmlElement("w:t")
    text.text = "DO NOT EXTRACT IMAGE TEXT"
    drawing.append(text)
    run._r.append(drawing)
    document.add_picture(io.BytesIO(png_bytes()))
    assert extract_document("answer.docx", docx_bytes(document))["text"] == "Real response"


def test_docx_image_only_has_empty_warning():
    document = Document()
    document.add_picture(io.BytesIO(png_bytes()))
    result = extract_document("scan.docx", docx_bytes(document))
    assert not result["text"]
    assert any("No extractable text was found" in warning for warning in result["warnings"])


def test_docx_corrupt_xml_is_meaningful():
    data = replace_zip_part(docx_bytes(Document()), "word/document.xml", b"<broken>")
    with pytest.raises(ValueError, match="Word document could not be read"):
        extract_document("broken.docx", data)


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_docx_entity_declarations_rejected(encoding):
    xml = '<?xml version="1.0"?><!DOCTYPE document [<!ENTITY x "unsafe">]><document/>'
    data = replace_zip_part(docx_bytes(Document()), "word/document.xml", xml.encode(encoding))
    with pytest.raises(ValueError, match="entity declarations"):
        extract_document("unsafe.docx", data)


def test_upload_size_limit(monkeypatch):
    monkeypatch.setattr(extract, "MAX_UPLOAD_BYTES", 5)
    with pytest.raises(ValueError, match="upload limit"):
        extract_document("essay.txt", b"123456")


@pytest.mark.parametrize("kind", ["txt", "pdf", "docx"])
def test_extracted_text_limit(monkeypatch, kind):
    monkeypatch.setattr(extract, "MAX_TEXT_CHARS", 5)
    document = Document()
    document.add_paragraph("123456")
    data = {"txt": b"123456", "pdf": pdf_bytes([("123456", False)]), "docx": docx_bytes(document)}[kind]
    with pytest.raises(ValueError, match="character limit"):
        extract_document("long." + kind, data)


def test_docx_expansion_limit(monkeypatch):
    monkeypatch.setattr(extract, "MAX_ZIP_EXPANDED_BYTES", 50)
    with pytest.raises(ValueError, match="safe expansion limits"):
        extract_document("large.docx", docx_bytes(Document()))


def test_docx_individual_xml_limit(monkeypatch):
    monkeypatch.setattr(extract, "MAX_XML_BYTES", 50)
    with pytest.raises(ValueError, match="XML part is too large"):
        extract_document("large.docx", docx_bytes(Document()))
