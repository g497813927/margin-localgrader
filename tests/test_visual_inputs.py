"""Generated charts and documents; no student images or model weights."""
import base64
import io
import json
from pathlib import Path
import subprocess
import sys
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, PngImagePlugin
from pypdf import PdfWriter, PdfReader
from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject, TextStringObject
from reportlab.pdfgen import canvas

from grader import visual_inputs as visual
from grader.app import app
from grader import vision


def image_bytes(size=(80, 40), color="red", format="PNG", **kwargs):
    output = io.BytesIO()
    Image.new("RGB", size, color).save(output, format, **kwargs)
    return output.getvalue()


def decoded(result):
    return Image.open(io.BytesIO(base64.b64decode(result["image_data_url"].split(",", 1)[1])))


def pdf_bytes():
    output = io.BytesIO()
    page = canvas.Canvas(output, pagesize=(300, 200))
    for color in [(1, 0, 0), (0, 0, 1)]:
        page.setFillColorRGB(*color)
        page.rect(50, 40, 150, 100, fill=1)
        page.drawString(20, 170, "Synthetic chart: values 1, 2, 3")
        page.showPage()
    page.save()
    return output.getvalue()


def docx_bytes():
    document = Document()
    document.add_paragraph("Synthetic diagram labels")
    document.add_picture(io.BytesIO(image_bytes(color="red")))
    document.add_picture(io.BytesIO(image_bytes(color="blue")))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def replace_part(data, name, replacement):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w") as target:
        for item in source.infolist():
            target.writestr(item.filename, replacement if item.filename == name else source.read(item))
    return output.getvalue()


@pytest.mark.parametrize("format,name", [("PNG", "chart.png"), ("JPEG", "chart.jpg"), ("WEBP", "chart.webp")])
def test_supported_images_and_no_metadata(format, name):
    result = visual.prepare_visual(name, image_bytes(format=format))
    assert result["count"] == result["index"] == 1
    assert result["width"] == 80 and result["height"] == 40
    assert decoded(result).format == "PNG"
    assert set(result) == {"image_data_url", "width", "height", "warnings", "index", "count", "source_label"}


def test_normalization_strips_metadata_and_bounds_dimensions():
    info = PngImagePlugin.PngInfo()
    info.add_text("Comment", "SYNTHETIC_PRIVATE_METADATA")
    result = visual.prepare_visual("large.png", image_bytes((6400, 2000), pnginfo=info))
    assert result["width"] == 4096 and result["height"] == 1280
    assert "SYNTHETIC_PRIVATE_METADATA".encode() not in base64.b64decode(result["image_data_url"].split(",")[1])
    assert result["warnings"]


def test_transparency_is_white_not_black():
    data = io.BytesIO()
    Image.new("RGBA", (10, 10), (0, 0, 0, 0)).save(data, "PNG")
    assert decoded(visual.prepare_visual("clear.png", data.getvalue())).getpixel((0, 0)) == (255, 255, 255)


def test_exif_orientation_applied_before_metadata_removed():
    exif = Image.Exif()
    exif[274] = 6
    result = visual.prepare_visual("rotated.jpg", image_bytes(format="JPEG", exif=exif))
    assert (result["width"], result["height"]) == (40, 80)
    assert not decoded(result).getexif()


def test_pixel_limit_checked_before_decode(monkeypatch):
    monkeypatch.setattr(visual, "MAX_PIXELS", 100)
    with pytest.raises(ValueError, match="million pixels"):
        visual.prepare_visual("large.png", image_bytes())


@pytest.mark.parametrize("name,data", [("bad.png", b"SYNTHETIC_PRIVATE_CORRUPT"), ("bad.svg", b"<svg/>"), ("bad.gif", b"GIF89a"), ("bad.png", b"")])
def test_invalid_images_are_errors_not_empty_evidence(name, data):
    with pytest.raises(ValueError) as error:
        visual.prepare_visual(name, data)
    assert "SYNTHETIC_PRIVATE_CORRUPT" not in str(error.value)


def test_animation_rejected():
    output = io.BytesIO()
    Image.new("RGB", (10, 10), "red").save(output, "PNG", save_all=True,
        append_images=[Image.new("RGB", (10, 10), "blue")])
    with pytest.raises(ValueError, match="Animated"):
        visual.prepare_visual("animated.png", output.getvalue())


@pytest.mark.parametrize("url", ["https://example.com/student.png", "file:///private/image.png", "data:text/html;base64,AAAA", "data:image/png;base64,%%%%"])
def test_urls_and_bad_encodings_never_fetched(url):
    with pytest.raises(ValueError):
        visual.normalize_image_data_url(url)


def test_vector_pdf_selected_page_rendered():
    first = visual.prepare_visual("chart.pdf", pdf_bytes())
    second = visual.prepare_visual("chart.pdf", pdf_bytes(), 2)
    assert first["count"] == second["count"] == 2
    assert first["image_data_url"] != second["image_data_url"]
    assert max(second["width"], second["height"]) <= 1600
    assert "page 2" in second["source_label"]


def test_pdf_lock_released_on_failure():
    with pytest.raises(ValueError):
        visual.prepare_visual("bad.pdf", b"SYNTHETIC_INVALID_PDF")
    assert not visual._PDF_LOCK.locked()
    assert visual.prepare_visual("good.pdf", pdf_bytes())["count"] == 2


def test_encrypted_pdf_rejected():
    output = io.BytesIO()
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf_bytes())))
    writer.encrypt("synthetic-password")
    writer.write(output)
    with pytest.raises(ValueError, match="unencrypted"):
        visual.prepare_visual("locked.pdf", output.getvalue())


def test_docx_embedded_images_in_reference_order():
    first = visual.prepare_visual("diagram.docx", docx_bytes())
    second = visual.prepare_visual("diagram.docx", docx_bytes(), 2)
    assert first["count"] == 2
    assert decoded(first).getpixel((0, 0)) == (255, 0, 0)
    assert decoded(second).getpixel((0, 0)) == (0, 0, 255)
    assert any("layout" in warning for warning in first["warnings"])


def test_word_external_images_and_xml_entities_not_read():
    data = docx_bytes()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        relationships = archive.read("word/_rels/document.xml.rels")
    altered = relationships.replace(b'Target="media/image1.png"', b'Target="https://example.com/private" TargetMode="External"')
    result = visual.prepare_visual("diagram.docx", replace_part(data, "word/_rels/document.xml.rels", altered))
    assert result["count"] == 1
    assert any("external" in warning for warning in result["warnings"])
    malicious = b'<!DOCTYPE x [<!ENTITY x SYSTEM "file:///private/test">]><x>&x;</x>'
    with pytest.raises(ValueError, match="entity"):
        visual.prepare_visual("diagram.docx", replace_part(data, "word/document.xml", malicious))


@pytest.mark.parametrize("index", [0, -1, 301, True, "1"])
def test_page_range_validation(index):
    with pytest.raises(ValueError):
        visual.prepare_visual("chart.png", image_bytes(), index)


def test_prepare_api_is_memory_only_and_cross_origin_rejected(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    client = TestClient(app)
    response = client.post("/api/vision/prepare", files={"file": ("test.png", image_bytes())})
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    assert list(tmp_path.iterdir()) == []
    assert client.post("/api/vision/prepare", headers={"Origin": "https://example.com"},
                       files={"file": ("test.png", image_bytes())}).status_code == 403


def test_analyze_api_has_no_grading_inputs_and_sanitizes_failures(monkeypatch):
    client = TestClient(app)
    image = visual.prepare_visual("test.png", image_bytes())["image_data_url"]
    calls = []
    def analyze(data, model, runtime_guard, **options):
        calls.append(model)
        raise vision.VisionUnavailable("Local vision is unavailable.")
    monkeypatch.setattr(vision, "analyze_image", analyze)
    payload = {"image_data_url": image, "model": "openbmb/minicpm-v4.7:1b"}
    result = client.post("/api/vision/analyze", json=payload)
    assert result.status_code == 503 and "scores" not in result.json()
    assert client.post("/api/vision/analyze", json={**payload, "rubric": {}}).status_code == 422
    assert len(calls) == 1


def pdf_with_image_metadata(dimensions, *, nested=False, external=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    images = DictionaryObject()
    for index, (width, height) in enumerate(dimensions):
        image = DecodedStreamObject()
        image.set_data(b"\0")
        image.update({NameObject("/Subtype"): NameObject("/Image"), NameObject("/Width"): NumberObject(width),
                      NameObject("/Height"): NumberObject(height), NameObject("/BitsPerComponent"): NumberObject(8),
                      NameObject("/ColorSpace"): NameObject("/DeviceGray")})
        if external:
            image[NameObject("/F")] = TextStringObject("https://example.invalid/private-image")
        images[NameObject(f"/I{index}")] = writer._add_object(image)
    resources = DictionaryObject({NameObject("/XObject"): images})
    if nested:
        form = DecodedStreamObject()
        form.set_data(b"")
        form.update({NameObject("/Subtype"): NameObject("/Form"), NameObject("/Resources"): resources})
        resources = DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/F0"): writer._add_object(form)})})
    page[NameObject("/Resources")] = resources
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.mark.parametrize("dimensions,nested", [([(5000, 5000)], False), ([(5000, 5000)], True),
                                               ([(4000, 4000)] * 3, True)])
def test_pdf_image_resource_limits_include_nested_forms_and_total_pixels(dimensions, nested):
    from grader.visual_worker import preflight_pdf, PdfPreparationError
    with pytest.raises(PdfPreparationError) as error:
        preflight_pdf(pdf_with_image_metadata(dimensions, nested=nested), 1)
    assert error.value.code == "resources"


def test_pdf_external_stream_rejected_before_native_decoder():
    from grader.visual_worker import preflight_pdf, PdfPreparationError
    with pytest.raises(PdfPreparationError) as error:
        preflight_pdf(pdf_with_image_metadata([(2, 2)], external=True), 1)
    assert error.value.code == "resources"


def test_pdf_inline_image_budget_checked_before_native_decoder():
    from grader.visual_worker import preflight_pdf, PdfPreparationError
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    stream = DecodedStreamObject()
    stream.set_data(b"BI /W 5000 /H 5000 /BPC 8 /CS /G ID x EI\n")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(PdfPreparationError) as error:
        preflight_pdf(output.getvalue(), 1)
    assert error.value.code == "resources"


def mock_pdf_worker(monkeypatch, *, output=b'{"error":"failed"}', timed_out=False, returncode=0):
    calls = []
    class Process:
        def __init__(self):
            self.returncode = None
            self.stdin = io.BytesIO()
            self.stdout = io.BytesIO()
            self.killed = False
            self.waited = False
        def communicate(self, *, input, timeout):
            calls.append((input, timeout))
            if timed_out:
                raise subprocess.TimeoutExpired("trusted-worker", timeout)
            self.returncode = returncode
            return output, None
        def poll(self): return self.returncode
        def kill(self):
            self.killed = True
            self.returncode = -9
        def wait(self, timeout): self.waited = True
    process = Process()
    def popen(argv, **kwargs):
        calls.append((argv, kwargs))
        return process
    monkeypatch.setattr(visual.subprocess, "Popen", popen)
    return process, calls


def test_pdf_worker_timeout_kills_reaps_and_releases_lock(monkeypatch):
    process, calls = mock_pdf_worker(monkeypatch, timed_out=True)
    with pytest.raises(visual.VisualInputError, match="30 seconds"):
        visual.prepare_visual("private-name.pdf", b"synthetic PDF")
    assert process.killed and process.waited
    assert process.stdin.closed and process.stdout.closed and not visual._PDF_LOCK.locked()
    argv, options = calls[0]
    assert argv == [sys.executable, "-I", "-B", str(Path(visual.__file__).with_name("visual_worker.py")), "1", "auto"]
    assert options["stderr"] == subprocess.DEVNULL and "shell" not in options
    assert "private-name" not in json.dumps(argv)
    assert "PYTHONPATH" not in options["env"] and "HTTPS_PROXY" not in options["env"]
    assert calls[1] == (b"synthetic PDF", 30)


@pytest.mark.parametrize("output,returncode", [(b'{"error":"PRIVATE_WORKER_DETAILS"}', 0),
                                               (b"PRIVATE_NATIVE_DIAGNOSTICS", 1), (b"invalid-json", 0),
                                               (b'{"image_data_url":"https://example.invalid/image","count":1}', 0)])
def test_pdf_worker_failure_never_echoes_native_or_external_input(monkeypatch, output, returncode):
    process, _ = mock_pdf_worker(monkeypatch, output=output, returncode=returncode)
    with pytest.raises(visual.VisualInputError) as error:
        visual.prepare_visual("synthetic.pdf", b"synthetic PDF")
    assert "PRIVATE" not in str(error.value) and "example.invalid" not in str(error.value)
    assert process.waited and not visual._PDF_LOCK.locked()


def test_pdf_worker_output_is_bounded(monkeypatch):
    monkeypatch.setattr(visual, "MAX_WORKER_OUTPUT_BYTES", 32)
    mock_pdf_worker(monkeypatch, output=b"x" * 33)
    with pytest.raises(visual.VisualInputError, match="rendered safely"):
        visual.prepare_visual("synthetic.pdf", b"synthetic PDF")


def test_docx_image_reference_count_matches_selectable_limit():
    data = docx_bytes()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        original = archive.read("word/document.xml")
    start = original.index(b"<w:drawing>")
    end = original.index(b"</w:drawing>", start) + len(b"</w:drawing>")
    amplified = original[:start] + original[start:end] * 301 + original[end:]
    with pytest.raises(visual.VisualInputError, match="300 image references"):
        visual.prepare_visual("diagrams.docx", replace_part(data, "word/document.xml", amplified))
